# Copyright (c) 2026, Vertex International Recruitment and contributors
# For license information, please see license.txt

"""Text extraction and normalization of lead identifiers (emails and WhatsApp numbers).

Pure parsing, no database access: everything that needs Frappe (country names, existing
leads) lives in vertex_outreach.contacts.
"""

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass
from typing import NamedTuple

import phonenumbers

# Loose on purpose: grab anything shaped like x@y so malformed addresses can be
# counted as invalid instead of silently disappearing.
CANDIDATE_RE = re.compile(r"[^\s<>;,\"'()\[\]{}:|/\\]+@[^\s<>;,\"'()\[\]{}:|/\\]+")

LOCAL_RE = re.compile(r"^[a-z0-9!#$%&*+=?^_`{}~-]+(\.[a-z0-9!#$%&*+=?^_`{}~-]+)*$")
DOMAIN_LABEL_RE = re.compile(r"^[a-z0-9]([a-z0-9-]{0,61}[a-z0-9])?$")
TLD_RE = re.compile(r"^[a-z]{2,63}$")

STRIP_CHARS = " \t\r\n.<>;,\"'()[]{}:"

# Outreach Contact.email is a unique varchar(140)
MAX_EMAIL_LENGTH = 140

GMAIL_DOMAINS = {"gmail.com", "googlemail.com"}

# ccTLDs that differ from the ISO 3166 code of their country
CCTLD_REGIONS = {"uk": "GB"}

# wa.me/256772123456, https://api.whatsapp.com/send?phone=256772123456&text=hi
WHATSAPP_LINK_RE = re.compile(
	r"(?:https?://)?(?:www\.)?(?:wa\.me/|api\.whatsapp\.com/send/?\?(?:[^\s#]*?&)?phone=)(?:\+|%2B)?(\d{6,15})",
	re.IGNORECASE,
)
# A run of digits with spaces, dashes and brackets, optionally starting with + or (+
PHONE_CANDIDATE_RE = re.compile(r"(?<![\w@+/.])(?:\+|\(\s*\+)?\s?\d[\d \t()\-]{5,22}\d(?![\w@])")
DATE_LIKE_RE = re.compile(r"^\(?(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4})")
PHONE_SEPARATORS_RE = re.compile(r"[\s().\-]")
# Free text is full of numbers; only runs this long look enough like a phone to report
MIN_TEXT_PHONE_DIGITS = 9
MIN_INTERNATIONAL_PHONE_DIGITS = 7
MAX_PHONE_DIGITS = 15

NEEDS_COUNTRY_CODE = "needs country code"
INVALID_PHONE = "not a valid phone number"
INVALID_EMAIL = "not a valid email address"

EMAIL = "Email"
WHATSAPP = "WhatsApp"

# Header cells are compared with everything but letters and digits removed: "E-mail" -> "email"
EMAIL_HEADERS = {"email", "emails", "emailaddress", "emailaddresses", "emailid", "mail"}
WHATSAPP_HEADERS = {
	"whatsapp",
	"whatsappnumber",
	"whatsappno",
	"whatsappphone",
	"wa",
	"wanumber",
	"phone",
	"phonenumber",
	"phoneno",
	"mobile",
	"mobilenumber",
	"mobileno",
	"cell",
	"cellphone",
	"telephone",
	"tel",
	"contactnumber",
}
ORGANIZATION_HEADERS = {
	"organization",
	"organisation",
	"company",
	"organizationname",
	"organisationname",
	"companyname",
	"business",
	"businessname",
	"agency",
}
COUNTRY_HEADERS = {"country", "countryname", "countrycode", "nation"}
HEADER_SCAN_ROWS = 20
MAX_ORGANIZATION_LENGTH = 140

# Words that label an identifier rather than name a company: "Tel:", "Email -"
LABEL_SUFFIX_RE = re.compile(
	r"(?:\b(?:e-?mail|mail|tel|telephone|phone|mobile|cell|whats\s?app|wa|contact|number|no)\b[\s.:#-]*)+$",
	re.IGNORECASE,
)
ORG_LEFT_STRIP = " \t,;|:-\u2013\u2014)>]\"'"
ORG_RIGHT_STRIP = " \t,;|:-\u2013\u2014(<[\"'"


class Phone(NamedTuple):
	e164: str | None
	region: str | None  # ISO 3166 alpha-2 of the number, e.g. "UG"
	error: str | None


@dataclass
class LeadInput:
	"""One lead found in pasted text or a file, or one identifier that could not be read."""

	input: str  # what the user typed, shown back to them
	organization: str | None = None
	email: str | None = None
	whatsapp: str | None = None  # E.164
	region: str | None = None  # ISO 3166 alpha-2 from the number
	country: str | None = None  # ISO 3166 alpha-2 from a country column
	error: str | None = None  # set when this is an unreadable identifier
	error_type: str | None = None  # EMAIL or WHATSAPP, for errors

	@property
	def type(self) -> str:
		if self.error:
			return self.error_type
		return " + ".join(t for t, v in ((EMAIL, self.email), (WHATSAPP, self.whatsapp)) if v)

	@property
	def regions(self) -> set[str]:
		"""Every country this lead points at: its number, its email's ccTLD and its country column."""
		return {r for r in (self.region, self.country, email_region(self.email) if self.email else None) if r}


def _extension(filename: str) -> str:
	return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _as_bytes(content: bytes | str) -> bytes:
	# File.get_content() returns str when the bytes happen to decode as utf-8
	return content.encode("utf-8") if isinstance(content, str) else content


def _as_text(content: bytes | str) -> str:
	if isinstance(content, str):
		return content.lstrip("﻿")
	return content.decode("utf-8-sig", errors="ignore")


def _cell(value) -> str:
	# Spreadsheets store 256772123456 as a float
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return "" if value is None else str(value).strip()


def extract_text(filename: str, content: bytes | str) -> str:
	"""Return all readable text from an uploaded file, chosen by extension."""
	ext = _extension(filename)
	if ext in ("pdf", "xlsx", "xlsm"):
		content = _as_bytes(content)

	if ext == "pdf":
		from pypdf import PdfReader

		reader = PdfReader(io.BytesIO(content))
		return "\n".join((page.extract_text() or "") for page in reader.pages)

	if ext in ("xlsx", "xlsm"):
		from openpyxl import load_workbook

		wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		try:
			return "\n".join(_row_text(row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
		finally:
			wb.close()

	return _as_text(content)


def _row_text(row) -> str:
	# " | " keeps neighbouring number cells from reading as one long number
	return " | ".join(_cell(v) for v in row if _cell(v))


def extract_tables(filename: str, content: bytes | str) -> list[list[tuple]] | None:
	"""Rows of each sheet for CSV/XLSX/XLSM files; None for formats without columns."""
	ext = _extension(filename)
	if ext == "csv":
		text = _as_text(content)
		try:
			dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
		except csv.Error:
			dialect = csv.excel
		return [[tuple(row) for row in csv.reader(io.StringIO(text), dialect)]]

	if ext in ("xlsx", "xlsm"):
		from openpyxl import load_workbook

		wb = load_workbook(io.BytesIO(_as_bytes(content)), read_only=True, data_only=True)
		try:
			return [list(ws.iter_rows(values_only=True)) for ws in wb.worksheets]
		finally:
			wb.close()
	return None


def parse_file(
	filename: str,
	content: bytes | str,
	default_country: str | None = None,
	resolve_country: Callable[[str], str | None] | None = None,
) -> list[LeadInput]:
	"""Leads in an uploaded file.

	For each CSV/XLSX sheet with a header row naming an email or WhatsApp column, each row is
	read as company / email / whatsapp / country. Everything else is read line by line with
	parse_text. `default_country` (ISO alpha-2) is used for numbers without a country code,
	after the row's own country column; `resolve_country` turns a country cell into ISO alpha-2.
	"""
	tables = extract_tables(filename, content)
	if tables is None:
		return parse_text(extract_text(filename, content), default_country)

	items = []
	for rows in tables:
		header = _find_header(rows)
		if header is None:
			for row in rows:
				items.extend(_parse_line(_row_text(row), default_country))
			continue
		header_row, columns = header
		for row in rows[header_row + 1 :]:
			items.extend(_parse_table_row(row, columns, default_country, resolve_country))
	return _dedupe(items)


def parse_text(text: str, default_country: str | None = None) -> list[LeadInput]:
	"""Leads in free text, one per line where a line holds at most one email and one number."""
	items = []
	for line in (text or "").splitlines():
		items.extend(_parse_line(line, default_country))
	return _dedupe(items)


def _dedupe(items: list[LeadInput]) -> list[LeadInput]:
	seen, unique = set(), []
	for item in items:
		key = (item.error_type, item.input.lower()) if item.error else (item.email, item.whatsapp)
		if key not in seen:
			seen.add(key)
			unique.append(item)
	return unique


def _parse_line(line: str, default_country: str | None) -> list[LeadInput]:
	found = []  # (position, LeadInput for one identifier)
	masked = list(line)

	for m in CANDIDATE_RE.finditer(line):
		masked[m.start() : m.end()] = " " * (m.end() - m.start())
		original = m.group().strip(STRIP_CHARS)
		email = normalize_email(original)
		found.append((m.start(), _identifier(original, email=email)))

	masked_line = "".join(masked)
	for m in WHATSAPP_LINK_RE.finditer(masked_line):
		masked[m.start() : m.end()] = " " * (m.end() - m.start())
		found.append((m.start(), _phone_identifier(m.group(), normalize_phone(m.group()))))

	for m in PHONE_CANDIDATE_RE.finditer("".join(masked)):
		original = m.group().strip()
		if not _looks_like_phone(original, MIN_TEXT_PHONE_DIGITS):
			continue
		found.append((m.start(), _phone_identifier(original, normalize_phone(original, default_country))))

	if not found:
		return []
	found.sort(key=lambda f: f[0])
	organization = _organization_before(line[: found[0][0]])
	return _combine([item for _pos, item in found], organization, line.strip())


def _parse_table_row(
	row: tuple, columns: dict[str, int], default_country: str | None, resolve_country
) -> list[LeadInput]:
	def cell(kind):
		col = columns.get(kind)
		return _cell(row[col]) if col is not None and col < len(row) else ""

	organization = cell("organization")[:MAX_ORGANIZATION_LENGTH] or None
	country = None
	if country_cell := cell("country"):
		if resolve_country:
			country = resolve_country(country_cell)
		elif len(country_cell) == 2 and country_cell.isalpha():
			country = country_cell.upper()

	found = [_identifier(original, email=email) for original, email in scan_emails(cell("email"))]
	whatsapp_cell = cell("whatsapp")
	links = [m.group() for m in WHATSAPP_LINK_RE.finditer(whatsapp_cell)]
	for part in links or re.split(r"[,;/\n|]+", whatsapp_cell):
		if part.strip() and any(c.isdigit() for c in part):
			found.append(_phone_identifier(part.strip(), normalize_phone(part, country or default_country)))

	items = _combine(found, organization, " | ".join(v for v in (cell(k) for k in columns) if v))
	for item in items:
		if not item.error:
			item.country = country
	return items


def _identifier(original: str, *, email: str | None) -> LeadInput:
	if email:
		return LeadInput(input=original, email=email)
	return LeadInput(input=original, error=INVALID_EMAIL, error_type=EMAIL)


def _phone_identifier(original: str, phone: Phone) -> LeadInput:
	if phone.e164:
		return LeadInput(input=original, whatsapp=phone.e164, region=phone.region)
	return LeadInput(input=original, error=phone.error, error_type=WHATSAPP)


def _combine(found: list[LeadInput], organization: str | None, whole: str) -> list[LeadInput]:
	"""One lead for the valid identifiers when there is at most one of each kind, else one each."""
	invalid = [f for f in found if f.error]
	emails = [f for f in found if f.email]
	phones = [f for f in found if f.whatsapp]
	if not emails and not phones:
		return invalid

	if len(emails) <= 1 and len(phones) <= 1:
		email, phone = (emails or [None])[0], (phones or [None])[0]
		single = email or phone
		lead = LeadInput(
			input=whole if (organization or (email and phone)) else single.input,
			organization=organization,
			email=email.email if email else None,
			whatsapp=phone.whatsapp if phone else None,
			region=phone.region if phone else None,
		)
		return [lead, *invalid]

	for f in emails + phones:
		f.organization = organization
	return emails + phones + invalid


def _organization_before(prefix: str) -> str | None:
	"""The company name on a line is whatever text comes before its first identifier."""
	text = prefix.strip(ORG_LEFT_STRIP).rstrip(ORG_RIGHT_STRIP)
	text = LABEL_SUFFIX_RE.sub("", text).rstrip(ORG_RIGHT_STRIP)
	if not re.search(r"[^\W\d_]", text):  # needs at least one letter
		return None
	return text[:MAX_ORGANIZATION_LENGTH]


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
	"""(row index, {kind: column}) of the first row naming an email or WhatsApp column, if any."""
	kinds = {
		"organization": ORGANIZATION_HEADERS,
		"email": EMAIL_HEADERS,
		"whatsapp": WHATSAPP_HEADERS,
		"country": COUNTRY_HEADERS,
	}
	for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
		keys = [re.sub(r"[^a-z0-9]", "", _cell(v).lower()) for v in row]
		columns = {}
		for kind, names in kinds.items():
			col = next((c for c, key in enumerate(keys) if key in names), None)
			if col is not None:
				columns[kind] = col
		if "email" in columns or "whatsapp" in columns:
			return i, columns
	return None


def find_partner_emails(filename: str, content: bytes | str) -> tuple[dict[str, str | None], int]:
	"""Emails with their organization: ({normalized email: organization or None}, invalid email count)."""
	partners: dict[str, str | None] = {}
	invalid = 0
	for item in parse_file(filename, content):
		if item.error:
			invalid += item.error_type == EMAIL
		elif item.email and (item.email not in partners or (item.organization and not partners[item.email])):
			partners[item.email] = item.organization
	return partners, invalid


def normalize_email(raw: str) -> str | None:
	"""Return the canonical form of an address, or None if it is not a valid email."""
	email = raw.strip().strip(STRIP_CHARS).lower()
	if email.startswith("mailto:"):
		email = email[7:]

	if email.count("@") != 1:
		return None
	local, domain = email.split("@")
	if not local or not domain or not LOCAL_RE.match(local):
		return None

	labels = domain.split(".")
	if len(labels) < 2 or not TLD_RE.match(labels[-1]):
		return None
	if not all(DOMAIN_LABEL_RE.match(label) for label in labels):
		return None

	if domain in GMAIL_DOMAINS:
		local = local.split("+", 1)[0].replace(".", "")
		domain = "gmail.com"
		if not local:
			return None

	email = f"{local}@{domain}"
	if len(email) > MAX_EMAIL_LENGTH:
		return None
	return email


def normalize_phone(raw, default_country: str | None = None) -> Phone:
	"""E.164 form and country of a phone number or WhatsApp link.

	Accepts spaces, dashes, brackets, a leading + or 00, and wa.me / api.whatsapp.com links.
	Numbers without a country code are read as national numbers of `default_country`
	(ISO alpha-2); without one they are invalid with NEEDS_COUNTRY_CODE.
	"""
	text = _cell(raw)
	if link := WHATSAPP_LINK_RE.search(text):
		number = "+" + link.group(1)  # WhatsApp links are always international
	else:
		number = PHONE_SEPARATORS_RE.sub("", text)
		if number.lower().startswith("tel:"):
			number = number[4:]
		if number.startswith("00"):
			number = "+" + number[2:]
	if not re.fullmatch(r"\+?\d+", number):
		return Phone(None, None, INVALID_PHONE)

	try:
		if number.startswith("+"):
			parsed = phonenumbers.parse(number, None)
		elif default_country:
			region = default_country.upper()
			parsed = phonenumbers.parse(number, region)
			# 256772123456 typed without the +
			calling_code = str(phonenumbers.country_code_for_region(region))
			if not phonenumbers.is_valid_number(parsed) and number.startswith(calling_code):
				parsed = phonenumbers.parse("+" + number, None)
		else:
			return Phone(None, None, NEEDS_COUNTRY_CODE)
	except phonenumbers.NumberParseException:
		return Phone(None, None, INVALID_PHONE)

	if not phonenumbers.is_valid_number(parsed):
		return Phone(None, None, INVALID_PHONE)
	region = phonenumbers.region_code_for_number(parsed)
	return Phone(
		phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164),
		region if region and region.isalpha() and len(region) == 2 else None,
		None,
	)


def _looks_like_phone(candidate: str, min_digits: int) -> bool:
	digits = sum(c.isdigit() for c in candidate)
	if candidate.startswith(("+", "(+")):
		min_digits = MIN_INTERNATIONAL_PHONE_DIGITS
	# + 2 for a leading 00
	return min_digits <= digits <= MAX_PHONE_DIGITS + 2 and not DATE_LIKE_RE.match(candidate)


def scan_emails(text: str) -> list[tuple[str, str | None]]:
	"""Every distinct email-like string in `text` as (original, normalized or None if invalid),
	in first-seen order."""
	seen: dict[str, str | None] = {}
	for match in CANDIDATE_RE.findall(text):
		original = match.strip(STRIP_CHARS)
		if original not in seen:
			seen[original] = normalize_email(original)
	return list(seen.items())


def find_emails(text: str) -> tuple[list[str], int]:
	"""Find addresses anywhere in `text`.

	Returns (unique normalized emails in first-seen order, count of distinct invalid candidates).
	"""
	valid: dict[str, None] = {}
	invalid: set[str] = set()
	for original, email in scan_emails(text):
		if email:
			valid[email] = None
		else:
			invalid.add(original.lower())
	return list(valid), len(invalid)


def email_domain(email: str) -> str:
	return email.rsplit("@", 1)[1]


def email_region(email: str) -> str | None:
	"""ISO alpha-2 country of an address's ccTLD (.ug -> UG, .uk -> GB), or None for .com etc."""
	tld = email.rsplit(".", 1)[-1].lower()
	if len(tld) != 2:
		return None
	return CCTLD_REGIONS.get(tld, tld.upper())
