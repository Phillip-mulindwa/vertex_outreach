# Copyright (c) 2026, Vertex International Recruitment and contributors
# For license information, please see license.txt

"""Text extraction, email discovery and normalization for outreach uploads."""

import io
import re

# Loose on purpose: grab anything shaped like x@y so malformed addresses can be
# counted as invalid instead of silently disappearing.
CANDIDATE_RE = re.compile(r"[^\s<>;,\"'()\[\]{}:|/\\]+@[^\s<>;,\"'()\[\]{}:|/\\]+")

LOCAL_RE = re.compile(r"^[a-z0-9!#$%&*+=?^_`{}~-]+(\.[a-z0-9!#$%&*+=?^_`{}~-]+)*$")
DOMAIN_LABEL_RE = re.compile(r"^[a-z0-9]([a-z0-9-]{0,61}[a-z0-9])?$")
TLD_RE = re.compile(r"^[a-z]{2,63}$")

STRIP_CHARS = " \t\r\n.<>;,\"'()[]{}:"

# Contact.name / email is a varchar(140) primary key.
MAX_EMAIL_LENGTH = 140

GMAIL_DOMAINS = {"gmail.com", "googlemail.com"}

_FREE_MAIL_BASES = (
	"gmail",
	"googlemail",
	"yahoo",
	"hotmail",
	"outlook",
	"live",
	"icloud",
	"aol",
	"proton",
	"protonmail",
	"btinternet",
	"sky",
	"msn",
)
FREE_MAIL_DOMAINS = frozenset(
	[f"{b}.com" for b in _FREE_MAIL_BASES] + [f"{b}.co.uk" for b in _FREE_MAIL_BASES] + ["proton.me"]
)


def extract_text(filename: str, content: bytes | str) -> str:
	"""Return all readable text from an uploaded file, chosen by extension."""
	ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

	if ext in ("pdf", "xlsx", "xlsm") and isinstance(content, str):
		# File.get_content() returns str when the bytes happen to decode as utf-8
		content = content.encode("utf-8")

	if ext == "pdf":
		from pypdf import PdfReader

		reader = PdfReader(io.BytesIO(content))
		return "\n".join((page.extract_text() or "") for page in reader.pages)

	if ext in ("xlsx", "xlsm"):
		from openpyxl import load_workbook

		wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		try:
			parts = []
			for ws in wb.worksheets:
				for row in ws.iter_rows(values_only=True):
					parts.append(" ".join(str(v) for v in row if v is not None))
			return "\n".join(parts)
		finally:
			wb.close()

	if isinstance(content, str):
		return content.lstrip("﻿")
	return content.decode("utf-8-sig", errors="ignore")


def normalize_email(raw: str) -> str | None:
	"""Return the canonical form of an address, or None if it is not a valid email."""
	email = raw.strip().strip(STRIP_CHARS).lower()
	if email.startswith("mailto:"):
		email = email[7:]

	if email.count("@") != 1:
		return None
	local, domain = email.split("@")
	if not local or not domain or not LOCAL_RE.match(local):
		return None

	labels = domain.split(".")
	if len(labels) < 2 or not TLD_RE.match(labels[-1]):
		return None
	if not all(DOMAIN_LABEL_RE.match(label) for label in labels):
		return None

	if domain in GMAIL_DOMAINS:
		local = local.split("+", 1)[0].replace(".", "")
		domain = "gmail.com"
		if not local:
			return None

	email = f"{local}@{domain}"
	if len(email) > MAX_EMAIL_LENGTH:
		return None
	return email


def scan_emails(text: str) -> list[tuple[str, str | None]]:
	"""Every distinct email-like string in `text` as (original, normalized or None if invalid),
	in first-seen order."""
	seen: dict[str, str | None] = {}
	for match in CANDIDATE_RE.findall(text):
		original = match.strip(STRIP_CHARS)
		if original not in seen:
			seen[original] = normalize_email(original)
	return list(seen.items())


def find_emails(text: str) -> tuple[list[str], int]:
	"""Find addresses anywhere in `text`.

	Returns (unique normalized emails in first-seen order, count of distinct invalid candidates).
	"""
	valid: dict[str, None] = {}
	invalid: set[str] = set()
	for original, email in scan_emails(text):
		if email:
			valid[email] = None
		else:
			invalid.add(original.lower())
	return list(valid), len(invalid)


def email_domain(email: str) -> str:
	return email.rsplit("@", 1)[1]
